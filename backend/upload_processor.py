# backend/upload_processor.py
# Enhanced file processing utilities with staged data architecture

import io
import re
import csv
import asyncio
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session
import json

# Use pandas for robust file processing
try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

# For Excel processing
try:
    import openpyxl
    import xlrd
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    
# Relative import
from . import models

class ProgressTracker:
    """Handles real-time progress tracking via WebSocket."""
    
    def __init__(self, session_id: str, websocket_manager=None):
        self.session_id = session_id
        self.websocket_manager = websocket_manager
        self.current_stage = "initializing"
        self.progress_percentage = 0.0
        self.rows_processed = 0
        self.total_rows = 0
        
    async def update_progress(self, stage: str, progress: float, rows_processed: int = 0, total_rows: int = 0, message: str = ""):
        """Send progress update via WebSocket."""
        self.current_stage = stage
        self.progress_percentage = min(progress, 100.0)
        self.rows_processed = rows_processed
        self.total_rows = total_rows
        
        if self.websocket_manager:
            await self.websocket_manager.send_progress(self.session_id, {
                "stage": stage,
                "progress": self.progress_percentage,
                "rows_processed": rows_processed,
                "total_rows": total_rows,
                "message": message,
                "timestamp": datetime.utcnow().isoformat()
            })
    
    async def send_final_update(self, result: dict):
        """Send final completion update."""
        if self.websocket_manager:
            await self.websocket_manager.send_progress(self.session_id, {
                "stage": "completed" if result.get('success') else "failed",
                "progress": 100.0,
                "final_result": result,
                "timestamp": datetime.utcnow().isoformat()
            })

class StagedTransactionProcessor:
    """Handles processing of uploaded transaction files with staged architecture."""
    
    def __init__(self, progress_tracker: Optional[ProgressTracker] = None):
        self.progress_tracker = progress_tracker
        
        # Known column patterns for auto-detection
        self.date_patterns = [
            r'date', r'transaction_date', r'datum', r'fecha', r'data', r'when'
        ]
        self.description_patterns = [
            r'description', r'beneficiary', r'memo', r'payee', r'vendor', 
            r'merchant', r'name', r'omschrijving', r'descripcion'
        ]
        self.amount_patterns = [
            r'amount', r'value', r'sum', r'bedrag', r'cantidad', r'monto', r'total'
        ]
        self.category_patterns = [
            r'category', r'type', r'classification', r'categorie', r'categoria'
        ]
        
    async def process_file_to_staged(
        self, 
        content: bytes, 
        filename: str, 
        file_type: str,
        upload_session: models.UploadSession,
        user: models.User,
        db: Session
    ) -> Dict[str, Any]:
        """Process file and create staged transactions."""
        
        if self.progress_tracker:
            await self.progress_tracker.update_progress(
                "reading_file", 5.0, 0, 0, "Reading file content..."
            )
        
        try:
            # Read and parse file
            if file_type == '.csv':
                raw_data = await self._process_csv_file(content)
            else:
                raw_data = await self._process_excel_file(content)
            
            if not raw_data:
                return {
                    "success": False,
                    "error": "No data found in file",
                    "total_rows": 0,
                    "staged_count": 0,
                    "error_count": 1,
                    "duplicate_count": 0,
                    "errors": ["File contains no readable data"]
                }
            
            # Store raw data in upload session
            upload_session.raw_data = raw_data[:100]  # Store first 100 rows for audit
            upload_session.total_rows = len(raw_data)
            db.commit()
            
            if self.progress_tracker:
                await self.progress_tracker.update_progress(
                    "parsing", 15.0, 0, len(raw_data), f"Parsed {len(raw_data)} rows"
                )
            
            # Detect format and normalize data
            format_detected = self._detect_format(raw_data)
            upload_session.format_detected = format_detected
            
            if self.progress_tracker:
                await self.progress_tracker.update_progress(
                    "normalizing", 25.0, 0, len(raw_data), f"Format detected: {format_detected}"
                )
            
            # Normalize transactions
            normalized_transactions = await self._normalize_transactions(raw_data, format_detected)
            
            if self.progress_tracker:
                await self.progress_tracker.update_progress(
                    "staging", 35.0, 0, len(normalized_transactions), "Creating staged transactions..."
                )
            
            # Stage transactions in batches
            result = await self._stage_transactions_batch(
                normalized_transactions, upload_session, user, db
            )
            
            result.update({
                "success": True,
                "total_rows": len(raw_data),
                "format_detected": format_detected
            })
            
            return result
            
        except Exception as e:
            error_msg = f"Processing failed: {str(e)}"
            
            if self.progress_tracker:
                await self.progress_tracker.update_progress(
                    "error", 0.0, 0, 0, error_msg
                )
            
            return {
                "success": False,
                "error": error_msg,
                "total_rows": 0,
                "staged_count": 0,
                "error_count": 1,
                "duplicate_count": 0,
                "errors": [error_msg]
            }
    
    async def _process_csv_file(self, content: bytes) -> List[Dict[str, Any]]:
        """Process CSV files with enhanced error handling."""
        try:
            # Try UTF-8 first, then fall back to other encodings
            encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']
            
            for encoding in encodings:
                try:
                    text_content = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError("Could not decode file with any supported encoding")
            
            # Use pandas if available for better CSV handling
            if PANDAS_AVAILABLE:
                try:
                    df = pd.read_csv(io.StringIO(text_content))
                    # Convert to list of dictionaries
                    return df.to_dict('records')
                except Exception:
                    # Fall back to manual processing
                    pass
            
            # Manual CSV processing
            csv_reader = csv.DictReader(io.StringIO(text_content))
            return list(csv_reader)
            
        except Exception as e:
            raise ValueError(f"Failed to parse CSV file: {str(e)}")
    
    async def _process_excel_file(self, content: bytes) -> List[Dict[str, Any]]:
        """Process Excel files with enhanced error handling."""
        if not EXCEL_AVAILABLE:
            raise ValueError("Excel processing not available. Install openpyxl and xlrd packages.")
        
        try:
            if PANDAS_AVAILABLE:
                # Use pandas for Excel processing
                df = pd.read_excel(io.BytesIO(content), engine='openpyxl')
                return df.to_dict('records')
            else:
                # Manual Excel processing with openpyxl
                from openpyxl import load_workbook
                
                workbook = load_workbook(io.BytesIO(content))
                worksheet = workbook.active
                
                # Get headers from first row
                headers = [cell.value for cell in worksheet[1]]
                
                # Process data rows
                data = []
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    row_data = dict(zip(headers, row))
                    # Skip empty rows
                    if any(value is not None for value in row_data.values()):
                        data.append(row_data)
                
                return data
                
        except Exception as e:
            raise ValueError(f"Failed to parse Excel file: {str(e)}")
    
    def _detect_format(self, raw_data: List[Dict[str, Any]]) -> str:
        """Detect the format/structure of the uploaded data."""
        if not raw_data:
            return "unknown"
        
        first_row = raw_data[0]
        columns = [str(key).lower() for key in first_row.keys()]
        
        # Check for common banking formats
        has_date = any(re.search('|'.join(self.date_patterns), col, re.IGNORECASE) for col in columns)
        has_description = any(re.search('|'.join(self.description_patterns), col, re.IGNORECASE) for col in columns)
        has_amount = any(re.search('|'.join(self.amount_patterns), col, re.IGNORECASE) for col in columns)
        
        if has_date and has_description and has_amount:
            # Try to identify specific bank formats
            if 'iban' in ' '.join(columns):
                return "dutch_bank"
            elif 'account' in ' '.join(columns):
                return "generic_bank"
            else:
                return "standard_transactions"
        
        return "unknown"
    
    async def _normalize_transactions(self, raw_data: List[Dict[str, Any]], format_detected: str) -> List[Dict[str, Any]]:
        """Normalize raw data into standard transaction format."""
        normalized = []
        
        for i, row in enumerate(raw_data):
            try:
                transaction = self._normalize_single_transaction(row, format_detected, i + 1)
                if transaction:
                    normalized.append(transaction)
            except Exception as e:
                # Skip invalid rows but log the error
                print(f"Warning: Skipping row {i + 1}: {str(e)}")
                continue
        
        return normalized
    
    def _normalize_single_transaction(self, row: Dict[str, Any], format_detected: str, row_num: int) -> Optional[Dict[str, Any]]:
        """Normalize a single transaction row."""
        transaction = {}
        
        # Find and parse date
        date_value = self._find_column_value(row, self.date_patterns)
        transaction['transaction_date'] = self._parse_date(date_value)
        
        # Find and parse description/beneficiary
        description_value = self._find_column_value(row, self.description_patterns)
        transaction['beneficiary'] = str(description_value) if description_value else ""
        
        # Find and parse amount
        amount_value = self._find_column_value(row, self.amount_patterns)
        transaction['amount'] = self._parse_amount(amount_value)
        
        # Find optional category
        category_value = self._find_column_value(row, self.category_patterns)
        transaction['category'] = str(category_value) if category_value else None
        
        # Set defaults
        transaction['labels'] = []
        transaction['is_private'] = False
        transaction['notes'] = None
        
        # Store original row data for audit
        transaction['raw_data'] = row
        
        # Validate required fields
        if not transaction['transaction_date']:
            raise ValueError(f"Missing or invalid date")
        if not transaction['beneficiary']:
            raise ValueError(f"Missing beneficiary/description")
        if transaction['amount'] is None:
            raise ValueError(f"Missing or invalid amount")
        
        return transaction
    
    def _find_column_value(self, row: Dict[str, Any], patterns: List[str]) -> Any:
        """Find value from row using column patterns."""
        for key, value in row.items():
            key_lower = str(key).lower()
            for pattern in patterns:
                if re.search(pattern, key_lower, re.IGNORECASE):
                    return value
        return None
    
    def _parse_date(self, date_value: Any) -> Optional[date]:
        """Parse date from various formats."""
        if not date_value:
            return None
        
        # If it's already a date object
        if isinstance(date_value, date):
            return date_value
        
        # If it's a datetime object
        if isinstance(date_value, datetime):
            return date_value.date()
        
        # Try to parse string dates
        date_str = str(date_value).strip()
        
        # Common date formats
        date_formats = [
            '%Y-%m-%d',
            '%d-%m-%Y',
            '%m/%d/%Y',
            '%d/%m/%Y',
            '%Y/%m/%d',
            '%Y-%m-%d %H:%M:%S',
            '%d-%m-%Y %H:%M:%S',
            '%Y%m%d'
        ]
        
        for fmt in date_formats:
            try:
                return datetime.strptime(date_str, fmt).date()
            except ValueError:
                continue
        
        return None
    
    def _parse_amount(self, amount_value: Any) -> Optional[Decimal]:
        """Parse amount from various formats."""
        if amount_value is None:
            return None
        
        # If it's already a number
        if isinstance(amount_value, (int, float, Decimal)):
            return Decimal(str(amount_value))
        
        # Clean string amount
        amount_str = str(amount_value).strip()
        
        # Remove common currency symbols and spaces
        amount_str = re.sub(r'[€$£¥₹\s]', '', amount_str)
        
        # Handle European decimal notation (comma as decimal separator)
        if ',' in amount_str and '.' in amount_str:
            # Both comma and dot - assume dot is thousands separator if it comes before comma
            if amount_str.rindex('.') < amount_str.rindex(','):
                amount_str = amount_str.replace('.', '').replace(',', '.')
        elif ',' in amount_str and amount_str.count(',') == 1:
            # Single comma - could be decimal separator
            if len(amount_str.split(',')[1]) <= 2:  # 2 digits after comma suggests decimal
                amount_str = amount_str.replace(',', '.')
        
        # Remove any remaining non-numeric characters except minus and dot
        amount_str = re.sub(r'[^-\d.]', '', amount_str)
        
        try:
            return Decimal(amount_str)
        except (InvalidOperation, ValueError):
            return None
    
    async def _stage_transactions_batch(
        self, 
        transactions: List[Dict[str, Any]], 
        upload_session: models.UploadSession,
        user: models.User,
        db: Session
    ) -> Dict[str, Any]:
        """Stage transactions in batches for better performance."""
        
        total_rows = len(transactions)
        batch_size = 50
        staged_count = 0
        error_count = 0
        duplicate_count = 0
        errors = []
        
        # Get existing transactions for duplicate detection
        existing_transactions = db.query(models.Transaction).filter(
            models.Transaction.owner_id == user.id
        ).all()
        
        existing_hashes = set()
        for tx in existing_transactions:
            hash_key = f"{tx.transaction_date}_{tx.beneficiary}_{tx.amount}"
            existing_hashes.add(hash_key)
        
        # Process in batches
        for batch_start in range(0, total_rows, batch_size):
            batch_end = min(batch_start + batch_size, total_rows)
            batch = transactions[batch_start:batch_end]
            
            if self.progress_tracker:
                progress = 35 + (batch_start / total_rows) * 55  # 35% to 90%
                await self.progress_tracker.update_progress(
                    "staging", 
                    progress, 
                    batch_start, 
                    total_rows,
                    f"Processing batch {batch_start // batch_size + 1}..."
                )
            
            # Process batch
            for idx, transaction in enumerate(batch):
                try:
                    # Check for duplicates
                    hash_key = f"{transaction['transaction_date']}_{transaction['beneficiary']}_{transaction['amount']}"
                    if hash_key in existing_hashes:
                        duplicate_count += 1
                        continue
                    
                    # Create staged transaction
                    staged_transaction = models.StagedTransaction(
                        transaction_date=transaction['transaction_date'],
                        beneficiary=transaction['beneficiary'],
                        amount=transaction['amount'],
                        category=transaction.get('category'),
                        labels=transaction.get('labels', []),
                        is_private=transaction.get('is_private', False),
                        notes=transaction.get('notes'),
                        user_id=user.id,
                        upload_session_id=upload_session.id,
                        status=models.TransactionStatus.STAGED,
                        confidence_score=Decimal('0.9'),  # High confidence for parsed data
                        raw_transaction_data=transaction.get('raw_data', {}),
                        created_at=datetime.utcnow()
                    )
                    
                    # Simple categorization suggestion
                    suggested_category = self._suggest_category(transaction['beneficiary'])
                    if suggested_category:
                        staged_transaction.suggested_category = suggested_category
                    
                    db.add(staged_transaction)
                    staged_count += 1
                    existing_hashes.add(hash_key)  # Prevent duplicates within same file
                    
                    # Commit in batches
                    if staged_count % 50 == 0:
                        db.commit()
                
                except Exception as e:
                    error_count += 1
                    errors.append(f"Row {batch_start + idx + 2}: {str(e)}")
            
            # Commit batch
            db.commit()
            
            # Add small delay to prevent overwhelming the system
            if batch_end < total_rows:
                await asyncio.sleep(0.1)
        
        # Final progress update
        if self.progress_tracker:
            await self.progress_tracker.update_progress(
                "finalizing", 
                95.0, 
                total_rows, 
                total_rows,
                "Finalizing staged transactions..."
            )
        
        # Update upload session
        upload_session.staged_count = staged_count
        upload_session.error_count = error_count
        upload_session.duplicate_count = duplicate_count
        upload_session.processed_rows = total_rows
        db.commit()
        
        return {
            "staged_count": staged_count,
            "error_count": error_count,
            "duplicate_count": duplicate_count,
            "errors": errors[:10]  # Limit errors in response
        }
    
    def _suggest_category(self, beneficiary: str) -> Optional[str]:
        """Simple category suggestion based on beneficiary name."""
        if not beneficiary:
            return None
        
        beneficiary_lower = beneficiary.lower()
        
        # Simple keyword-based categorization
        category_keywords = {
            "Groceries": ["supermarket", "grocery", "food", "albert heijn", "jumbo", "lidl", "aldi"],
            "Transport": ["uber", "taxi", "train", "bus", "fuel", "gas", "petrol", "parking"],
            "Utilities": ["electricity", "gas", "water", "internet", "phone", "mobile"],
            "Entertainment": ["cinema", "netflix", "spotify", "restaurant", "bar", "pub"],
            "Shopping": ["amazon", "shop", "store", "mall", "clothing", "fashion"],
            "Healthcare": ["doctor", "pharmacy", "hospital", "medical", "dentist"],
            "Banking": ["bank", "atm", "fee", "charge", "interest"]
        }
        
        for category, keywords in category_keywords.items():
            if any(keyword in beneficiary_lower for keyword in keywords):
                return category
        
        return None
    
    def _validate_transaction_data(self, transaction: Dict[str, Any], row_num: int) -> Optional[str]:
        """Validate transaction data and return error message if invalid."""
        if not transaction.get('transaction_date'):
            return f"Row {row_num}: Missing or invalid date"
        
        if not transaction.get('beneficiary'):
            return f"Row {row_num}: Missing beneficiary/description"
        
        if transaction.get('amount') is None:
            return f"Row {row_num}: Missing or invalid amount"
        
        return None
    
    async def validate_file_structure(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Validate file structure without full processing."""
        try:
            file_type = '.' + filename.split('.')[-1].lower()
            
            if file_type == '.csv':
                raw_data = await self._process_csv_file(content)
            else:
                raw_data = await self._process_excel_file(content)
            
            if not raw_data:
                return {"valid": False, "error": "No data found in file"}
            
            format_detected = self._detect_format(raw_data)
            issues = []
            suggestions = []
            
            # Check for required columns based on detected format
            first_row = raw_data[0]
            columns = list(first_row.keys())
            
            if format_detected == "unknown":
                issues.append("Could not detect a known format")
                suggestions.append("Ensure file has date, description, and amount columns")
            
            # Check first few rows for data quality
            sample_size = min(5, len(raw_data))
            for i in range(sample_size):
                try:
                    self._normalize_single_transaction(raw_data[i], format_detected, i + 1)
                except Exception as e:
                    issues.append(f"Row {i + 2}: {str(e)}")
            
            return {
                "valid": len(issues) == 0,
                "format_detected": format_detected,
                "estimated_rows": len(raw_data),
                "columns_found": columns,
                "issues": issues,
                "suggestions": suggestions
            }
            
        except Exception as e:
            return {"valid": False, "error": str(e)}